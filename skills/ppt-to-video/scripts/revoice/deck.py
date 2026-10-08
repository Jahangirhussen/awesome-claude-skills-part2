"""PPTX 인제스트 — 회사가 준 발표자료를 '그대로' 렌더해서 씬 소재로 쓴다.

첫 제작 시도의 실패 원인이 슬라이드를 HTML로 재디자인한 것이었으므로
(브랜드 룩 불일치·빈 여백·⚠확인 마커 노출), 여기서는 디자인을 만들지 않는다.
PowerPoint COM으로 슬라이드를 1920x1080 PNG로 내보내고, 본문 텍스트와
발표자 노트만 뽑아 나레이션 기획의 입력으로 쓴다.
"""
import os
import json
import subprocess

W, H = 1920, 1080


def _ps(script: str) -> str:
    """PowerShell 실행 — pywin32 없이 COM을 쓰기 위한 얇은 우회."""
    r = subprocess.run(["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
                       capture_output=True, text=True)
    if r.returncode != 0:
        raise RuntimeError(f"powershell failed: {r.stderr.strip()}")
    return r.stdout


def export_slides(pptx: str, out_dir: str) -> list:
    """PPTX의 각 슬라이드를 out_dir/slide_NN.png 로 내보내고 경로 목록 반환.

    PowerPoint가 원본 폰트·사진·도형을 그대로 렌더하므로 브랜드 룩이 보존된다.
    """
    pptx = os.path.abspath(pptx)
    out_dir = os.path.abspath(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    _ps(f"""
$app = New-Object -ComObject PowerPoint.Application
$pres = $app.Presentations.Open('{pptx}', $true, $false, $false)
foreach ($i in 1..$pres.Slides.Count) {{
  $pres.Slides.Item($i).Export(('{out_dir}\\slide_{{0:d2}}.png' -f $i), 'PNG', {W}, {H})
}}
$pres.Close(); $app.Quit()
""")
    return sorted(os.path.join(out_dir, f) for f in os.listdir(out_dir)
                  if f.startswith("slide_") and f.endswith(".png"))


def extract_outline(pptx: str) -> list:
    """슬라이드별 {n, title, body[], notes} — 나레이션 기획의 원자료."""
    from pptx import Presentation
    out = []
    for n, slide in enumerate(Presentation(pptx).slides, 1):
        lines = [ln.strip() for sh in slide.shapes if sh.has_text_frame
                 for ln in sh.text_frame.text.splitlines() if ln.strip()]
        notes = ""
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
        out.append({"n": n, "title": lines[0] if lines else "",
                    "body": lines[1:], "notes": notes})
    return out


def ingest(pptx: str, out_dir: str = "deck") -> dict:
    """슬라이드 PNG + outline.json 을 만들고 요약을 반환."""
    pngs = export_slides(pptx, out_dir)
    outline = extract_outline(pptx)
    for item in outline:                       # 슬라이드 번호로 PNG 경로 연결
        p = os.path.join(out_dir, f"slide_{item['n']:02d}.png")
        item["png"] = p.replace("\\", "/") if os.path.exists(p) else ""
    path = os.path.join(out_dir, "outline.json")
    with open(path, "w", encoding="utf-8") as f:
        json.dump({"source_pptx": pptx, "slides": outline}, f, ensure_ascii=False, indent=2)
    return {"pngs": len(pngs), "outline": path, "slides": len(outline)}


def render_animated_slides(pptx: str, durations: dict, out_dir: str = "deck/anim",
                           effect_duration: float = 0.5, stagger: float = 0.25,
                           width: int = W, fps: int = 30) -> dict:
    """슬라이드별 요소 순차 페이드인 애니메이션 mp4 생성 — {slide_no: 노출초} → {slide_no: mp4경로}.

    정지 이미지에 눈에 안 띄는 푸시인만 얹던 기존 방식 대신, PowerPoint 자체 애니메이션
    엔진으로 요소가 순서대로 나타나는 모션을 만든다. HTML 슬라이드로 새로 그리지 않고
    원본 pptx의 폰트·사진·레이아웃을 PowerPoint 렌더러가 그대로 써서 브랜드 룩이 보존된다.
    노출초는 보통 그 씬 나레이션 TTS 길이 — 짧으면 페이드가 안 끝난 채 다음 씬으로 넘어가고
    길면 마지막 요소가 뜬 채로 정지하니, 호출 전에 리드인을 더한 최종 길이를 넘겨야 한다.

    durations의 각 값은 초(float) 하나이거나, 나레이션 문장별 등장 타이밍을 지정하는
    `{"duration": 초, "checkpoints": [문장1이 끝나는 시각, 문장2가 끝나는 시각, ...]}`
    형태일 수 있다 — 후자는 화면 요소를 문서 순서대로 그 구간에 배분해, 나레이터가
    실제로 그 문장을 말하는 시점에 맞춰 등장하게 한다(build.py의 _sentence_checkpoints
    가 실제 TTS 음성을 전사해 계산).
    """
    pptx = os.path.abspath(pptx)
    out_dir = os.path.abspath(out_dir)
    os.makedirs(out_dir, exist_ok=True)
    durations_path = os.path.join(out_dir, "_durations.json")
    with open(durations_path, "w", encoding="utf-8") as f:
        json.dump({str(k): v for k, v in durations.items()}, f)

    script = os.path.join(os.path.dirname(__file__), "_animate_slides.ps1")
    timeout = 60 + 45 * len(durations)   # 슬라이드당 인코딩+오버헤드 여유
    r = subprocess.run(
        ["powershell", "-NoProfile", "-NonInteractive", "-File", script,
         "-PptxPath", pptx, "-DurationsJson", durations_path, "-OutDir", out_dir,
         "-EffectDuration", str(effect_duration), "-Stagger", str(stagger),
         "-Width", str(width), "-Fps", str(fps)],
        capture_output=True, text=True, timeout=timeout)
    if r.returncode != 0:
        raise RuntimeError(f"슬라이드 애니메이션 렌더 실패: {r.stderr.strip() or r.stdout.strip()}")

    out = {}
    for line in r.stdout.splitlines():
        parts = line.split(maxsplit=2)
        if len(parts) == 3 and parts[0] == "RENDERED":
            out[int(parts[1])] = parts[2].strip()
        elif parts and parts[0] == "FAILED":
            print(f"[deck] 슬라이드 {parts[1]} 애니메이션 렌더 실패(재시도 포함) — "
                  f"정지 이미지로 대체됩니다.")
        elif parts and parts[0] == "RETRY":
            print(f"[deck] {' '.join(parts[1:])} — 재시도 중")
    return out


def main(argv=None):
    import sys
    argv = argv if argv is not None else sys.argv[1:]
    if not argv:
        print("usage: python -m revoice.deck <파일.pptx> [출력디렉터리]")
        raise SystemExit(1)
    r = ingest(argv[0], argv[1] if len(argv) > 1 else "deck")
    print(f"[deck] 슬라이드 {r['slides']}장 · PNG {r['pngs']}장 → {r['outline']}")


if __name__ == "__main__":
    main()
